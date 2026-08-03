import os
import re
import argparse
import logging
from typing import Optional, Tuple, List, Dict, Any

import pdfplumber
import docx
from devlens.data_collection.github_client import GitHubClient

logger = logging.getLogger(__name__)

# Reserved path words on github.com that are not usernames
GITHUB_RESERVED_WORDS = {
    "topics", "features", "settings", "pricing", "about", "marketplace",
    "search", "orgs", "organizations", "login", "join", "explore",
    "trending", "security", "enterprise", "customer-stories", "readme",
    "sponsors", "collections", "events", "issues", "pulls", "notifications"
}


class ResumeExtractor:
    """Extracts and verifies GitHub username from a PDF or DOCX resume."""

    def __init__(self, client: Optional[GitHubClient] = None):
        self.client = client or GitHubClient()

    def extract_text_from_pdf(self, file_path: str) -> str:
        """Extract text from all pages of a PDF file using pdfplumber."""
        text_chunks = []
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_chunks.append(page_text)
        return "\n".join(text_chunks)

    def extract_text_from_docx(self, file_path: str) -> str:
        """Extract text from paragraphs and tables of a DOCX file using python-docx."""
        doc = docx.Document(file_path)
        text_chunks = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_chunks.append(paragraph.text.strip())
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    if cell.text.strip():
                        text_chunks.append(cell.text.strip())
        return "\n".join(text_chunks)

    def extract_raw_text(self, file_path: str) -> str:
        """Determine file type and extract raw text content."""
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"Resume file not found at: {file_path}")

        ext = os.path.splitext(file_path)[1].lower()
        if ext == ".pdf":
            return self.extract_text_from_pdf(file_path)
        elif ext in (".docx", ".doc"):
            return self.extract_text_from_docx(file_path)
        else:
            raise ValueError(f"Unsupported file format: '{ext}'. Supported formats: .pdf, .docx")

    def _find_direct_url_username(self, text: str) -> Optional[str]:
        """Strategy 1: Search for direct github.com/<username> URL pattern in text."""
        # Matches patterns like github.com/username or https://github.com/username
        pattern = r'(?:https?://)?(?:www\.)?github\.com/([a-zA-Z0-9\-]+)(?:[/\s\?\)\"\']|$)'
        matches = re.findall(pattern, text, re.IGNORECASE)
        for match in matches:
            candidate = match.strip().rstrip('/')
            if candidate.lower() not in GITHUB_RESERVED_WORDS and len(candidate) > 0:
                return candidate
        return None

    def _find_labeled_username(self, text: str) -> Optional[str]:
        """Strategy 2: Search for 'GitHub:' or 'GitHub Profile:' labeled fields."""
        patterns = [
            r'github\s*(?:profile|handle|account|user)?\s*:\s*@?([a-zA-Z0-9\-]+)',
            r'gh\s*:\s*@?([a-zA-Z0-9\-]+)',
        ]
        for pattern in patterns:
            matches = re.findall(pattern, text, re.IGNORECASE)
            for match in matches:
                candidate = match.strip()
                if candidate.lower() not in GITHUB_RESERVED_WORDS and len(candidate) > 0:
                    # Ignore generic placeholder words
                    if candidate.lower() in ("http", "https", "url", "link", "profile", "user"):
                        continue
                    return candidate
        return None

    def _infer_and_verify_username(self, text: str) -> Optional[str]:
        """Strategy 3: Fall back to inferring candidate usernames from email or name, verified via REST API."""
        candidates = []

        # 1. Extract email addresses
        email_pattern = r'([a-zA-Z0-9._%+-]+)@[a-zA-Z0-9.-]+\.[a-zA-Z]{2,}'
        email_matches = re.findall(email_pattern, text)
        for prefix in email_matches:
            # Generate variations from email prefix (e.g. john.doe -> johndoe, john-doe)
            clean_prefix = re.sub(r'[^a-zA-Z0-9\-._]', '', prefix)
            candidates.append(clean_prefix)
            candidates.append(clean_prefix.replace('.', ''))
            candidates.append(clean_prefix.replace('.', '-'))

        # 2. Extract potential candidate from first line (typically Full Name)
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        if lines:
            first_line = lines[0]
            # Simple name candidate (e.g. "John Doe" -> "johndoe", "john-doe")
            name_parts = re.findall(r'[a-zA-Z]+', first_line)
            if 1 < len(name_parts) <= 3:
                full_name_concat = "".join(name_parts).lower()
                full_name_hyphen = "-".join(name_parts).lower()
                candidates.append(full_name_concat)
                candidates.append(full_name_hyphen)

        # Deduplicate candidates while preserving order
        unique_candidates = []
        for cand in candidates:
            cand_clean = cand.strip().strip('-').strip('.')
            if cand_clean and cand_clean.lower() not in GITHUB_RESERVED_WORDS:
                if cand_clean not in unique_candidates:
                    unique_candidates.append(cand_clean)

        logger.info(f"Inference candidates to verify via REST API: {unique_candidates}")

        # Verify candidates against GitHub REST API
        for candidate in unique_candidates:
            res = self.client.rest_request(f"users/{candidate}")
            if res and isinstance(res, dict) and "login" in res:
                logger.info(f"Verified candidate '{candidate}' via GitHub API.")
                return res["login"]

        return None

    def extract_github_username(self, file_path: str) -> Dict[str, Any]:
        """Extract GitHub username from resume using prioritized strategies."""
        text = self.extract_raw_text(file_path)

        # Strategy 1: Direct GitHub URL
        username = self._find_direct_url_username(text)
        if username:
            return {
                "username": username,
                "confidence": "direct_url",
                "details": f"Found direct GitHub URL match in resume text: '{username}'",
                "file_path": file_path,
            }

        # Strategy 2: Labeled Field
        username = self._find_labeled_username(text)
        if username:
            return {
                "username": username,
                "confidence": "labeled_field",
                "details": f"Found labeled GitHub handle match in resume text: '{username}'",
                "file_path": file_path,
            }

        # Strategy 3: Inference & REST verification
        username = self._infer_and_verify_username(text)
        if username:
            return {
                "username": username,
                "confidence": "inferred_verified",
                "details": f"Inferred from resume contact info and verified via GitHub REST API: '{username}'",
                "file_path": file_path,
            }

        return {
            "username": None,
            "confidence": "none",
            "details": "No GitHub profile URL, handle label, or verifiable inferred username found in resume.",
            "file_path": file_path,
        }


def main():
    parser = argparse.ArgumentParser(description="Extract GitHub username from a PDF or DOCX resume.")
    parser.add_argument("file_path", type=str, help="Path to resume file (.pdf or .docx)")
    args = parser.parse_args()

    extractor = ResumeExtractor()
    try:
        result = extractor.extract_github_username(args.file_path)
        print("\n" + "=" * 55)
        print("  DEV LENS - RESUME GITHUB EXTRACTOR")
        print("=" * 55)
        print(f" File Path   : {result['file_path']}")
        print(f" Username    : {result['username'] if result['username'] else 'Not Found (None)'}")
        print(f" Confidence  : {result['confidence']}")
        print(f" Details     : {result['details']}")
        print("=" * 55 + "\n")
    except Exception as e:
        logger.error(f"Extraction failed: {e}")
        raise


if __name__ == "__main__":
    main()
